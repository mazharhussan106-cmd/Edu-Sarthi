# Batch 6 builder: the last words needed to reach exactly 2,000.
# Owns: which files make up Batch 6 (g01–g05 from the HTML list, m01–m04 written in full, and 45 rows
# from the 910-word workbook with the 2 new columns from z910.py) and its Read Me.
# Does not own: row layout and checks (sheet.py), Batches 1–5, or the merged 2,000-word file (final.py).
#
#   VOCAB_JSON="10000000 new/vocabulary_database.json" WM910="10000000 new/WordMaster1000_Words_1001-plus.xlsx" \
#     python3 build6.py out.xlsx WordMaster2000_{Sample,Batch1,Batch2,Batch3,Batch4,Batch5}_*.xlsx
import json, os, sys
from openpyxl import load_workbook
sys.path.insert(0, sys.path[0])
import g01, g02, g03, g04, g05, m01, m02, m03, m04, z910
from removed import REMOVED, REMOVED_B6, NOT_ADDED_B6
from fixes import FIX
from sheet import fields, h_row, n_row, old_words, check, write_book, headers_from

OUT, OLD = sys.argv[1], sys.argv[2:]
START = 1808
VOCAB = {w["word"].lower(): w for w in json.load(open(os.environ["VOCAB_JSON"]))["words"]}

rows, overridden = [], []
for mod in (g01, g02, g03, g04, g05):
    O = getattr(mod, "O", {})
    for f in fields(mod, 35):
        o, v = O.get(f[0], {}), VOCAB[f[1].lower()]
        rows.append(h_row(f, o, v))
        if "mw" in o or "mc" in o:
            cm = v["common_mistake"]
            overridden.append([f[0], f[1], cm["wrong"], cm["correct"], o.get("mw", cm["wrong"]), o.get("mc", cm["correct"])])
n_h = len(rows)
for mod in (m01, m02, m03, m04):
    rows += [n_row(f) for f in fields(mod, 52)]
n_n = len(rows) - n_h

# The 910-word workbook rows keep their 54 columns; z910 adds the spelling tip and action sequence.
tips = {}
for line in z910.TEXT.strip().splitlines():
    wid, tip, act = [p.strip() for p in (line + " |").split("|")[:3]]
    tips[wid] = (tip, [a.strip() for a in act.split(";")] if act else None)
wm = {r[1]: r for r in list(load_workbook(os.environ["WM910"], read_only=True)["Words 1001+"].iter_rows(values_only=True))[2:] if r[1]}
for wid, (tip, act) in tips.items():
    r = list(wm[wid])
    a = f"Before: {act[0]} → During: {act[1]} → After: {act[2]}" if act else ""
    rows.append(r[1:20] + [a] + r[20:26] + [tip] + r[26:54])
rows = [[START + i] + r for i, r in enumerate(rows)]
HEADERS = headers_from(OLD[-2])
for r in rows:
    for header, value in FIX.get(r[1], {}).items():
        r[HEADERS.index(header)] = value

old = old_words(OLD)
removed = {**REMOVED, **REMOVED_B6}
assert len(old) == 1807 and all(i in old for i in removed), len(old)
check(rows, old, removed)
assert 1550 - len(removed) + (1807 - 1550) + len(rows) == 2000, len(rows)

END = START + len(rows) - 1
notes = [
    (f"WordMaster 2000 — Batch 6: words {START}–{END} (the last batch)", True),
    (f"Contents: {n_h} words from the HTML word list, {n_n} words written in full for this file, and "
     f"{len(rows) - n_h - n_n} words from the 910-word workbook (those keep their original 54 columns; only the "
     "2 new columns were added).", False),
    ("How the list was checked: every word was matched exactly (not by a loose rule) against words 1–1807. "
     "Repeats were dropped — for example 'be able to' (Batch 5 has 'able'), 'hurry' (Batch 5 has 'hurry up'), "
     "'firstly' (this batch has 'first of all'). Weaker words were swapped for core ones that were still missing: "
     "give up, get on, work on, look up, carry on, would like, look like, feel like, get ready, pay attention, "
     "let me know, at the moment, in the end, by mistake, on purpose, home, married, noisy, silly, awesome, "
     "pardon and happy birthday. The full list with reasons is on the 'Not added' sheet.", False),
    (f"Total: 1,550 − {len(removed)} removed + 257 (Batch 5) + {len(rows)} here = 2,000. "
     "'give up on' joins the Removed list because 'give up' is now in.", False),
    ("", False),
    ("Where each field came from:", True),
    ("• HTML-list words: part of speech, category, IPA, CEFR, frequency, Hindi meaning, simple explanation, "
     "Example 1, collocations, grammar pattern, common mistake, memory trick, word family, synonyms, antonyms and "
     "the fill-in-the-blank come from vocabulary_database.json. Everything else was written for this file.", False),
    ("• Where the HTML list's mistake line was unclear, it was replaced — see 'Corrections'.", False),
    ("• Words written in full: every field was written for this file.", False),
    ("• 910-workbook words: all 54 original columns are unchanged; Spelling Tip and Action Sequence were added.", False),
    ("• AI-estimated, not from a corpus: CEFR level, frequency, importance.", False),
    ("• Real-life lines are illustrative. They are not real quotes from any film, news outlet or person.", False),
    ("", False),
    ("S.No continues from 1808. WordMaster2000_FINAL.xlsx has all 2,000 words numbered 1–2,000.", False),
]
sheets = [
    ("Removed", ["Word ID", "Word", "Why it was taken out"], [[i, old[i], why] for i, why in REMOVED_B6.items()], [12, 22, 60]),
    ("Not added", ["Planned word", "Why it was not added"], [[w, why] for w, why in NOT_ADDED_B6.items()], [24, 60]),
    ("Corrections", ["Word ID", "Word", "HTML list's mistake", "HTML list's correction", "New mistake", "New correction"],
     overridden, [12, 14, 40, 40, 40, 40]),
]
write_book(OUT, rows, HEADERS, sheets, notes, START)
print("saved", OUT, len(rows), "rows,", len(overridden), "corrections")
