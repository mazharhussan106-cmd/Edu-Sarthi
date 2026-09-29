# Final merge: all batches → one workbook of exactly 2,000 words, numbered 1–2,000, plus the Extra list.
# Owns: dropping the removed words, renumbering, the coverage summary, and the Removed, Not added and
# Corrections sheets.
# Does not own: any word content — every row is copied unchanged from the batch files except S.No.
#
#   FREQ_LIST=en_50k.txt python3 final.py WordMaster2000_FINAL.xlsx \
#     WordMaster2000_{Sample,Batch1,Batch2,Batch3,Batch4,Batch5,Batch6,Extra}_*.xlsx
import os, sys
from collections import Counter
from openpyxl import load_workbook
sys.path.insert(0, sys.path[0])
from removed import REMOVED, REMOVED_B6, NOT_ADDED, NOT_ADDED_B6, RESTORED_TO_EXTRA
from sheet import write_book, headers_from
from coverage import coverage

OUT = sys.argv[1]
BATCHES = [p for p in sys.argv[2:] if "_Extra_" not in p]
EXTRA = [p for p in sys.argv[2:] if "_Extra_" in p]
removed = {**REMOVED, **REMOVED_B6}

rows, words_by_id, corrections = [], {}, []
for path in BATCHES:
    wb = load_workbook(path, read_only=True)
    batch = path.split("_")[1]
    for r in list(wb.worksheets[1].iter_rows(values_only=True))[2:]:
        if not r[1]:
            continue
        words_by_id[r[1]] = r[2]
        if r[1] not in removed:
            rows.append(list(r))
    if "Corrections" in wb.sheetnames:
        # Batches 1–4 corrected the app's examples; Batches 5–6 corrected the HTML list's. Same 6 columns.
        source = "HTML word list" if batch in ("Batch5", "Batch6") else "Speak Sarthi app"
        for c in list(wb["Corrections"].iter_rows(values_only=True))[1:]:
            if c[0] and c[0] not in removed:
                corrections.append([batch, source] + list(c))

assert len(rows) == 2000, len(rows)
ids, words = [r[1] for r in rows], [str(r[2]).lower() for r in rows]
assert len(set(ids)) == 2000 and len(set(words)) == 2000, "duplicate ID or word"
assert all(i in words_by_id for i in removed)
for i, r in enumerate(rows, start=1):
    r[0] = i
    assert all(v not in ("", None) for k, v in enumerate(r) if k != 20), (r[1], "blank cell")
    assert r[3] not in ("Verb", "Phrasal Verb") or r[20], (r[1], "verb without an action sequence")
    assert str(r[49]).split(" → ")[0].strip() != str(r[23]).strip(), (r[1], "find & fix repeats the mistake")

extra = [list(r) for p in EXTRA for r in list(load_workbook(p, read_only=True).worksheets[1].iter_rows(values_only=True))[2:] if r[1]]
# 'mean' is in both on purpose: ADJ-366 = unkind, VRB-739 = signify.
assert not ({str(r[2]).lower() for r in extra} & set(words)) - {"mean"}, "Extra repeats a main-list word"

pos = Counter(str(r[3]).split(" / ")[0] for r in rows)
cefr = Counter(str(r[8]) for r in rows)
notes = [
    ("WordMaster 2000 — FINAL (2,000 words)", True),
    ("What this is: the 2,000 most useful non-noun words and phrases for spoken English, in the 56-column format "
     "(the 54-column WordMaster reference + Action Sequence and Spelling Tip). Nouns are not included, as agreed.", False),
    ("How it was built: the six batch files, merged in order, with the words on the 'Removed' sheet taken out, "
     "then numbered 1–2,000. No word content was changed in the merge.", False),
    (f"Removed: {len(removed)} words that repeated another entry (toward/towards, somebody/someone …) or were "
     "too rare for a top-2,000 list (thirteen–nineteen, rare ordinals, furthermore …).", False),
    ("Not added: words that were planned but dropped as repeats or as less useful than the core words added "
     "instead — see the 'Not added' sheet.", False),
    ("Corrections: every place where a source's mistake example was replaced, with the original and the new line.", False),
    ("", False),
    ("Mix of the list:", True),
    ("By part of speech: " + ", ".join(f"{k} {v}" for k, v in pos.most_common()), False),
    ("By CEFR level (AI-estimated): " + ", ".join(f"{k} {v}" for k, v in sorted(cefr.items())), False),
    ("", False),
    ("Good to know:", True),
    ("• CEFR level, frequency and importance are AI estimates, not from a corpus.", False),
    ("• Real-life lines are illustrative. They are not real quotes from any film, news outlet or person.", False),
    ("• Word IDs are kept from the source files (VRB-…, PHR-…, AUX-…), so rows can be matched back to the app.", False),
    ("• Rows are in batch order, not ranked. Use the filter on 'Importance' or 'CEFR Level' to study the most "
     "important or easiest words first.", False),
]
if extra:
    notes += [
        ("", False),
        (f"Extra sheet ({len(extra)} words, numbered {extra[0][0]}–{extra[-1][0]}):", True),
        ("Words beyond the 2,000 that make the list cover every teachable word in the 1,000 most common words of "
         "spoken English: non-noun gaps (yeah, need, live, going to, gonna, uh/um …), five words moved back from "
         "the Removed list (somebody, everybody, nobody, anybody, till) and the most common nouns (people, day, "
         "money, phone …). Names, swear words and subtitle marks like 'sighs' are left out on purpose.", False),
    ]
if extra and os.environ.get("FREQ_LIST"):
    h = headers_from(BATCHES[-1])
    main, both = coverage(rows, h.index("Word Forms"), os.environ["FREQ_LIST"]), coverage(rows + extra, h.index("Word Forms"), os.environ["FREQ_LIST"])
    notes += [
        ("", False),
        ("Coverage of spoken English (OpenSubtitles 2018 — film and TV dialogue, top 50,000 words):", True),
        (f"• 2,000 words: {100 * main['all']:.1f}% of all spoken words. With Extra: {100 * both['all']:.1f}%.", False),
        (f"• Top 100 words: {100 * main[100]:.1f}% → {100 * both[100]:.1f}%. Top 500: {100 * main[500]:.1f}% → "
         f"{100 * both[500]:.1f}%. Top 1,000: {100 * main[1000]:.1f}% → {100 * both[1000]:.1f}%.", False),
        (f"• Still not covered in the top 1,000: {len(both['missing_1000'])} words, all names and titles (Mr, Dr, "
         "John, America), swear words, or subtitle marks ('sighs', single letters) — none belong in a learner's "
         "word list. Every other word in the top 1,000 is covered.", False),
        ("• Coverage counts a word's forms (go/went/going, don't = do + not) and the words inside phrases "
         "('course' in 'of course').", False),
    ]
sheets = [
    ("Removed", ["Word ID", "Word", "Why it was taken out"],
     [[i, words_by_id[i], why + (" — moved to the Extra sheet" if i in RESTORED_TO_EXTRA else "")]
      for i, why in removed.items()], [12, 22, 60]),
    ("Not added", ["Planned word", "Why it was not added"],
     [[w, "less useful than the core words added instead"] for w in NOT_ADDED] + [[w, why] for w, why in NOT_ADDED_B6.items()],
     [24, 60]),
    ("Corrections", ["Batch", "Source", "Word ID", "Word", "Original mistake", "Original correction", "New mistake", "New correction"],
     corrections, [10, 18, 12, 14, 36, 36, 36, 36]),
]
write_book(OUT, rows, headers_from(BATCHES[-1]), sheets, notes, 1, extra)
print("saved", OUT, len(rows), "words +", len(extra), "extra,", len(removed), "removed,", len(corrections), "corrections")
print(pos.most_common(), sorted(cefr.items()))
