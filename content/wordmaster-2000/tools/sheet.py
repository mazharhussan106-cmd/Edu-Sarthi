# Shared by build5.py and build6.py: turning H-format (HTML-list) and N-format (hand-written) lines
# into 56-column rows, checking them, and writing the workbook in the WordMaster layout.
# Does not own: which words go in a batch, or the Read Me text — each batch builder decides those.
import re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

PLAN = "Day 1 learn → Day 3 → Day 7 → Day 14 → Day 30 review"
SECTIONS = [
    ("1. RECOGNITION", "3F2A8C", 12), ("2. MEANING & USAGE", "0E7C7B", 9), ("3. GRAMMAR & ERRORS", "B5452B", 7),
    ("4. MEMORY & WORD-BUILDING", "C07A10", 7), ("5. CONVERSATION & COMPARISON", "2E6DB4", 7),
    ("6. REAL-LIFE USAGE (illustrative, not real quotes)", "6B7280", 3), ("7. PRACTICE ZONE", "1E8E5A", 7),
    ("8. REVIEW / SPACED REPETITION", "7C3AED", 4),
]
arial = lambda **k: Font(name="Arial", **k)
thin = Side(style="thin", color="D9D9D9")


def fields(mod, n):
    for line in mod.TEXT.strip().splitlines():
        f = [p.strip() for p in line.split("|")]
        assert len(f) == n, (mod.__name__, f[0], len(f))
        yield f


def action(s):
    if s in ("-", "—", ""):
        return ""
    b, d, a = [p.strip() for p in s.split(";")]
    return f"Before: {b} → During: {d} → After: {a}"


def options(s):
    opts = [o.strip() for o in s.split(";")]
    assert len(opts) == 4, s
    return "  ".join(f"{'ABCD'[k]}) {o}" for k, o in enumerate(opts))


def recall(meaning):
    m = meaning.rstrip(".")
    return f"Which word means '{m[0].lower() + m[1:]}'?"


def or_dash(xs, what):
    return ", ".join(xs) or f"—  (no common {what})"


def h_row(f, o, v):
    """One H line (35 fields) + its vocabulary_database.json entry v + overrides o → 55 cells (no S.No)."""
    word = f[1]
    cm = v["common_mistake"]
    # The HTML list writes its fill-in-the-blank as "Fill in: '…'"; the answer column adds "→ answer".
    fb = o.get("fb") or re.sub(r"^Fill in:\s*['\"](.*)['\"]$", r"\1", v["practice"]["question"]).replace("\\'", "'") \
        + " → " + v["practice"]["answer"]
    meaning = o.get("meaning", v["meaning_en"])
    return [
        f[0], word, o.get("pos", v["part_of_speech"]), v["category_name"], v["ipa"], f[2], f[3], v["cefr"],
        v["frequency"], int(f[4]), f[5],
        o.get("hm", v["meaning_hi"]), meaning[0].upper() + meaning[1:] + ".", o.get("ex1", v["example"]), f[6], f[7],
        o.get("col", ", ".join(v["collocations"])), f[8], f[9], action(f[10]),
        o.get("gp", v["grammar_pattern"]), f[11], o.get("mw", cm["wrong"]), o.get("mc", cm["correct"]), o.get("why", f[12]), f[13], f[14],
        f[15], f[16], o.get("fam", ", ".join(v["word_family"]) or "—"), f[17], o.get("mt", v["memory_trick"]), f[18], f[19],
        o.get("syn", or_dash(v["synonyms"], "synonym")), o.get("ant", or_dash(v["antonyms"], "antonym")), f[20], f[21], f[22], f[23], f[24],
        f[25], f[26], f[27],
        f[28], options(f[29]), f[30], fb, f[31], f[32], f[33],
        PLAN, recall(meaning), word, f[34],
    ]


def n_row(f):
    """One N line (52 fields, every column hand-written) → 55 cells (no S.No)."""
    return [
        f[0], f[1], f[2], f[3], f[4], f[5], f[6], f[7], f[8], int(f[9]), f[10],
        f[11], f[12], f[13], f[14], f[15], f[16], f[17], f[18], action(f[19]),
        f[20], f[21], f[22], f[23], f[24], f[25], f[26],
        f[27], f[28], f[29], f[30], f[31], f[32], f[33],
        f[34], f[35], f[36], f[37], f[38], f[39], f[40],
        f[41], f[42], f[43],
        f[44], options(f[45]), f[46], f[47], f[48], f[49], f[50],
        PLAN, recall(f[12]), f[1], f[51],
    ]


def old_words(paths):
    """Word ID → word for every row already published in these workbooks."""
    old = {}
    for path in paths:
        for r in list(load_workbook(path, read_only=True).worksheets[1].iter_rows(values_only=True))[2:]:
            if r[1]:
                old[r[1]] = str(r[2])
    return old


def check(rows, old, removed, strict_mistakes=True, same_word_ok=()):
    # same_word_ok: headwords allowed twice because they are different words (mean = unkind / mean = signify).
    kept = {w.lower() for i, w in old.items() if i not in removed} - set(same_word_ok)
    ids, words = [r[1] for r in rows], [r[2].lower() for r in rows]
    assert len(set(ids)) == len(ids) and not set(ids) & set(old), "duplicate word ID"
    assert len(set(words)) == len(words), "duplicate word in batch"
    assert not set(words) & kept, sorted(set(words) & kept)
    for r in rows:
        assert len(r) == 56, r[1]
        blank = [k + 1 for k, v in enumerate(r) if v in ("", None) and k != 20]
        assert not blank, (r[1], blank)
        assert r[47] in ("A", "B", "C", "D"), (r[1], r[47])
        # Helping and modal verbs are grammar words with no before/during/after, like adjectives.
        assert bool(r[20]) == (r[3] in ("Verb", "Phrasal Verb")), (r[1], "action", r[3])
        if strict_mistakes:
            # Find & Fix should practise a different error from the Common Mistake column.
            assert str(r[49]).split(" → ")[0].strip() != str(r[23]).strip(), (r[1], "find & fix repeats the mistake")


def write_book(out, rows, headers, sheets, notes, start, extra=()):
    wb = Workbook()
    rm = wb.active
    rm.title = "Read Me"
    words_sheet(wb, f"Words {start}–{start + len(rows) - 1}", rows, headers)
    if extra:
        words_sheet(wb, f"Extra {extra[0][0]}–{extra[-1][0]}", extra, headers)
    for title, hdr, data, widths in sheets:
        s = wb.create_sheet(title)
        for j, h in enumerate(hdr, start=1):
            c = s.cell(1, j, h)
            c.font, c.fill = arial(sz=10, b=True, color="FFFFFF"), PatternFill("solid", fgColor="3F2A8C")
            s.column_dimensions[get_column_letter(j)].width = widths[j - 1]
        for i, r in enumerate(data, start=2):
            for j, v in enumerate(r, start=1):
                c = s.cell(i, j, v)
                c.font, c.alignment = arial(sz=10), Alignment(wrap_text=True, vertical="top")
        s.freeze_panes = "A2"

    rm.column_dimensions["A"].width = 120
    for r, (t, bold) in enumerate(notes, start=1):
        c = rm.cell(r, 1, t)
        c.font, c.alignment = arial(sz=12 if r == 1 else 10, b=bold), Alignment(wrap_text=True, vertical="top")
    wb.save(out)


def words_sheet(wb, title, rows, headers):
    ws = wb.create_sheet(title)
    col = 1
    for title, colour, n in SECTIONS:
        fill = PatternFill("solid", fgColor=colour)
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + n - 1)
        c = ws.cell(1, col, title)
        c.font, c.fill, c.alignment = arial(sz=11, b=True, color="FFFFFF"), fill, Alignment(vertical="center")
        for k in range(n):
            ws.cell(1, col + k).fill = fill
            h = headers[col + k - 1]
            c = ws.cell(2, col + k, h)
            c.font = arial(sz=10, b=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="E8590C") if "★NEW" in h else fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
            if "★NEW" in h:
                c.comment = Comment("New column, not in the 54-column reference. Added from the "
                                    "Premium Flashcard Master Specification.", "Claude")
        col += n
    for i, r in enumerate(rows, start=3):
        for j, v in enumerate(r, start=1):
            c = ws.cell(i, j, v)
            c.font, c.alignment, c.border = arial(sz=10), Alignment(wrap_text=True, vertical="top"), Border(bottom=thin)
        ws.row_dimensions[i].height = 150
    ws.row_dimensions[2].height = 42
    ws.column_dimensions["A"].width = 7
    for j in range(2, 57):
        ws.column_dimensions[get_column_letter(j)].width = 34 if j in (14, 21, 25, 26, 38, 42) else 22
    ws.freeze_panes = "D3"
    ws.auto_filter.ref = f"A2:{get_column_letter(56)}{len(rows) + 2}"


def headers_from(path):
    # Same 56 headers as build.py, read from a published batch so the builders can't drift apart.
    h = list(list(load_workbook(path, read_only=True).worksheets[1].iter_rows(values_only=True))[1])
    assert len(h) == 56 and sum(n for _, _, n in SECTIONS) == 56
    return h
